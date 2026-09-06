import os
import json
import uuid
import logging
import asyncio
import hashlib
import re
from typing import List, Optional, Dict, Any
from datetime import datetime
from pathlib import Path

from app.core.config import settings
from app.models.wiki import wiki_store, WikiPage

logger = logging.getLogger(__name__)


class DocumentImportService:
    def __init__(self):
        self._import_dir = os.path.join(settings.DATA_DIR, "imports")
        self._paddle_ocr_instance = None  # PaddleOCR 实例缓存
        os.makedirs(self._import_dir, exist_ok=True)

    async def import_pdf(self, file_path: str, space_id: str = "default",
                         author: str = "system") -> Optional[WikiPage]:
        try:
            import fitz
            doc = fitz.open(file_path)
            content_parts = []
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text()
                if text.strip():
                    content_parts.append(f"## 第{page_num + 1}页\n\n{text}")
            doc.close()
            content = "\n\n".join(content_parts)
            title = Path(file_path).stem
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"PDF导入失败: {e}")
            return None

    async def import_word(self, file_path: str, space_id: str = "default",
                          author: str = "system") -> Optional[WikiPage]:
        try:
            from docx import Document
            doc = Document(file_path)
            content_parts = []
            for para in doc.paragraphs:
                if para.style.name.startswith("Heading"):
                    level = para.style.name.replace("Heading ", "").strip()
                    try:
                        level = int(level)
                    except ValueError:
                        level = 1
                    content_parts.append(f"{'#' * level} {para.text}")
                else:
                    content_parts.append(para.text)

            for table in doc.tables:
                table_md = "| " + " | ".join([cell.text for cell in table.rows[0].cells]) + " |\n"
                table_md += "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |\n"
                for row in table.rows[1:]:
                    table_md += "| " + " | ".join([cell.text for cell in row.cells]) + " |\n"
                content_parts.append(table_md)

            content = "\n\n".join(content_parts)
            title = Path(file_path).stem
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"Word导入失败: {e}")
            return None

    async def import_markdown(self, file_path: str, space_id: str = "default",
                              author: str = "system") -> Optional[WikiPage]:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
            title = Path(file_path).stem
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"Markdown导入失败: {e}")
            return None

    async def import_excel(self, file_path: str, space_id: str = "default",
                           author: str = "system") -> Optional[WikiPage]:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True)
            content_parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                content_parts.append(f"## {sheet_name}\n")
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    header = rows[0]
                    content_parts.append("| " + " | ".join([str(c or "") for c in header]) + " |")
                    content_parts.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for row in rows[1:]:
                        content_parts.append("| " + " | ".join([str(c or "") for c in row]) + " |")
            wb.close()
            content = "\n\n".join(content_parts)
            title = Path(file_path).stem
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"Excel导入失败: {e}")
            return None

    async def import_excel_enhanced(self, file_path: str, space_id: str = "default",
                                    author: str = "system",
                                    include_stats: bool = True,
                                    max_rows_per_sheet: int = 1000) -> Optional[WikiPage]:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True)
            content_parts = []
            content_parts.append(f"# {Path(file_path).stem}\n")
            content_parts.append(f"**源文件**: {Path(file_path).name}\n")

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                content_parts.append(f"\n## 工作表: {sheet_name}\n")
                rows = list(ws.iter_rows(values_only=True, max_row=max_rows_per_sheet))
                if not rows:
                    content_parts.append("（空工作表）\n")
                    continue

                header = rows[0]
                content_parts.append("### 数据表格\n")
                content_parts.append("| " + " | ".join([str(c or "") for c in header]) + " |")
                content_parts.append("| " + " | ".join(["---"] * len(header)) + " |")
                for row in rows[1:]:
                    content_parts.append("| " + " | ".join([str(c or "") for c in row]) + " |")

                if include_stats and len(rows) > 1:
                    content_parts.append("\n### 数据统计\n")
                    content_parts.append(f"- 总行数: {len(rows) - 1}")
                    content_parts.append(f"- 总列数: {len(header)}")
                    numeric_cols = {}
                    for col_idx, col_name in enumerate(header):
                        col_name_str = str(col_name or f"列{col_idx + 1}")
                        values = []
                        for row in rows[1:]:
                            val = row[col_idx] if col_idx < len(row) else None
                            if val is not None:
                                try:
                                    values.append(float(val))
                                except (ValueError, TypeError):
                                    pass
                        if values:
                            numeric_cols[col_name_str] = values

                    for col_name, values in numeric_cols.items():
                        content_parts.append(f"\n**{col_name}**:")
                        content_parts.append(f"  - 最小值: {min(values):.2f}")
                        content_parts.append(f"  - 最大值: {max(values):.2f}")
                        content_parts.append(f"  - 平均值: {sum(values) / len(values):.2f}")
                        content_parts.append(f"  - 数据量: {len(values)}")

            wb.close()
            content = "\n".join(content_parts)
            title = Path(file_path).stem
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"增强Excel导入失败: {e}")
            return None

    async def import_excel_to_structured(self, file_path: str, space_id: str = "default",
                                         author: str = "system") -> Dict[str, Any]:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True)
            sheets_data = {}
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    headers = [str(c or "") for c in rows[0]]
                    records = []
                    for row in rows[1:]:
                        record = {}
                        for idx, header in enumerate(headers):
                            record[header] = str(row[idx]) if idx < len(row) and row[idx] is not None else ""
                        records.append(record)
                    sheets_data[sheet_name] = {"headers": headers, "records": records, "total": len(records)}
            wb.close()
            headers_str = ", ".join(sheets_data.keys())
            content = f"# {Path(file_path).stem} (结构化)\n\n工作表: {headers_str}\n\n"
            for sname, sdata in sheets_data.items():
                content += f"## {sname} ({sdata['total']}条记录)\n\n"
                h = sdata["headers"]
                content += "| " + " | ".join(h) + " |\n"
                content += "| " + " | ".join(["---"] * len(h)) + " |\n"
                for rec in sdata["records"][:50]:
                    content += "| " + " | ".join([rec.get(hh, "") for hh in h]) + " |\n"
                if sdata["total"] > 50:
                    content += f"\n> 仅显示前50条，共{sdata['total']}条记录\n"
                content += "\n"
            page = wiki_store.create_page(
                title=Path(file_path).stem, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
            return {"success": True, "page_id": page.page_id, "sheets": sheets_data}
        except Exception as e:
            logger.error(f"结构化Excel导入失败: {e}")
            return {"success": False, "error": str(e)}

    async def import_image_ocr(self, file_path: str, space_id: str = "default",
                               author: str = "system") -> Optional[WikiPage]:
        try:
            from PIL import Image
            img = Image.open(file_path)
            text = await self._ocr_image(img)
            if not text.strip():
                text = "（图片中未识别到文字内容）"
            content = f"## OCR识别结果\n\n{text}"
            title = Path(file_path).stem
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"图片OCR导入失败: {e}")
            return None

    async def _ocr_image(self, image) -> str:
        if not settings.local_ocr_enabled:
            return ""
        try:
            return await self._ocr_paddle(image)
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"PaddleOCR识别失败，尝试Tesseract: {e}")
        try:
            import pytesseract
            return pytesseract.image_to_string(image, lang="chi_sim+eng")
        except ImportError:
            logger.warning("pytesseract未安装，OCR功能不可用")
            return ""
        except Exception as e:
            logger.error(f"OCR识别失败: {e}")
            return ""

    async def _ocr_paddle(self, image) -> str:
        if not settings.local_ocr_enabled:
            return ""
        if self._paddle_ocr_instance is None:
            from paddleocr import PaddleOCR
            self._paddle_ocr_instance = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
        ocr = self._paddle_ocr_instance
        import numpy as np
        if hasattr(image, 'save'):
            import io
            buf = io.BytesIO()
            image.save(buf, format='PNG')
            buf.seek(0)
            from PIL import Image
            img_array = np.array(Image.open(buf))
        else:
            img_array = np.array(image)
        result = ocr.ocr(img_array, cls=True)
        lines = []
        if result and result[0]:
            for line in result[0]:
                if line and len(line) >= 2:
                    text = line[1][0] if isinstance(line[1], (list, tuple)) else str(line[1])
                    lines.append(text)
        return "\n".join(lines)

    def _preprocess_image(self, image):
        try:
            from PIL import ImageEnhance, ImageFilter
            if image.mode != 'L':
                image = image.convert('L')
            enhancer = ImageEnhance.Contrast(image)
            image = enhancer.enhance(2.0)
            enhancer = ImageEnhance.Sharpness(image)
            image = enhancer.enhance(1.5)
            image = image.filter(ImageFilter.MedianFilter(size=3))
            return image
        except Exception as e:
            logger.warning(f"图片预处理失败: {e}")
            return image

    async def import_image_ocr_enhanced(self, file_path: str, space_id: str = "default",
                                        author: str = "system",
                                        preprocess: bool = True,
                                        lang: str = "chi_sim+eng") -> Optional[WikiPage]:
        try:
            from PIL import Image
            img = Image.open(file_path)
            if preprocess:
                img = self._preprocess_image(img)
            text = await self._ocr_image(img)
            if not text.strip():
                text = "（图片中未识别到文字内容）"
            content_parts = [f"## OCR识别结果\n"]
            content_parts.append(f"**源文件**: {Path(file_path).name}\n")
            content_parts.append(f"**预处理**: {'是' if preprocess else '否'}\n")
            content_parts.append(f"\n### 识别文本\n\n{text}")
            content = "\n".join(content_parts)
            title = Path(file_path).stem + "_OCR"
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"增强OCR导入失败: {e}")
            return None

    async def batch_ocr(self, file_paths: List[str], space_id: str = "default",
                        author: str = "system", preprocess: bool = True) -> Dict[str, Any]:
        results = {"success": [], "failed": []}
        for fp in file_paths:
            ext = Path(fp).suffix.lower()
            if ext not in (".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp"):
                results["failed"].append({"file": fp, "error": f"非图片格式: {ext}"})
                continue
            page = await self.import_image_ocr_enhanced(fp, space_id, author, preprocess)
            if page:
                results["success"].append({"file": fp, "page_id": page.page_id, "title": page.title})
            else:
                results["failed"].append({"file": fp, "error": "OCR识别失败"})
        return results

    async def import_webpage(self, url: str, space_id: str = "default",
                             author: str = "system") -> Optional[WikiPage]:
        try:
            import httpx
            from bs4 import BeautifulSoup
            async with httpx.AsyncClient(timeout=30, follow_redirects=True) as client:
                resp = await client.get(url)
                resp.raise_for_status()
            soup = BeautifulSoup(resp.text, "lxml")
            for tag in soup(["script", "style", "nav", "footer", "header"]):
                tag.decompose()
            title = soup.title.string.strip() if soup.title and soup.title.string else url
            body = soup.find("body") or soup
            text = body.get_text(separator="\n", strip=True)
            paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
            content = f"## {title}\n\n来源: {url}\n\n" + "\n\n".join(paragraphs)
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"网页抓取失败: {e}")
            return None

    async def import_notion(self, export_dir: str, space_id: str = "default",
                            author: str = "system") -> List[WikiPage]:
        pages = []
        export_path = Path(export_dir)
        if not export_path.exists():
            logger.error(f"Notion导出目录不存在: {export_dir}")
            return pages
        for md_file in export_path.rglob("*.md"):
            page = await self.import_markdown(str(md_file), space_id, author)
            if page:
                pages.append(page)
        for csv_file in export_path.rglob("*.csv"):
            page = await self._import_csv(str(csv_file), space_id, author)
            if page:
                pages.append(page)
        logger.info(f"Notion导入完成: {len(pages)} 个页面")
        return pages

    async def import_notion_api(self, database_id: str, notion_token: str,
                                space_id: str = "default",
                                author: str = "system") -> Dict[str, Any]:
        results = {"success": [], "failed": []}
        try:
            import httpx
            headers = {
                "Authorization": f"Bearer {notion_token}",
                "Notion-Version": "2022-06-28",
                "Content-Type": "application/json",
            }
            async with httpx.AsyncClient(timeout=30) as client:
                url = f"https://api.notion.com/v1/databases/{database_id}/query"
                resp = await client.post(url, headers=headers, json={})
                resp.raise_for_status()
                data = resp.json()

                for page_data in data.get("results", []):
                    try:
                        page_id = page_data.get("id", "")
                        title_prop = ""
                        props = page_data.get("properties", {})
                        for prop_name, prop_val in props.items():
                            if prop_val.get("type") == "title":
                                title_arr = prop_val.get("title", [])
                                if title_arr:
                                    title_prop = "".join(t.get("plain_text", "") for t in title_arr)
                                break

                        block_url = f"https://api.notion.com/v1/blocks/{page_id}/children"
                        block_resp = await client.get(block_url, headers=headers)
                        block_resp.raise_for_status()
                        blocks = block_resp.json().get("results", [])

                        content_parts = []
                        for block in blocks:
                            block_type = block.get("type", "")
                            text_content = self._extract_notion_block_text(block, block_type)
                            if text_content:
                                if block_type.startswith("heading_"):
                                    level = block_type.replace("heading_", "")
                                    try:
                                        level = int(level)
                                    except ValueError:
                                        level = 2
                                    content_parts.append(f"{'#' * min(level, 4)} {text_content}")
                                elif block_type == "bulleted_list_item":
                                    content_parts.append(f"- {text_content}")
                                elif block_type == "numbered_list_item":
                                    content_parts.append(f"1. {text_content}")
                                elif block_type == "quote":
                                    content_parts.append(f"> {text_content}")
                                elif block_type == "code":
                                    lang = block.get("code", {}).get("language", "")
                                    content_parts.append(f"```{lang}\n{text_content}\n```")
                                else:
                                    content_parts.append(text_content)

                        title = title_prop or f"Notion页面_{page_id[:8]}"
                        content = "\n\n".join(content_parts) or "（空页面）"
                        wiki_page = wiki_store.create_page(
                            title=title, content=content, space_id=space_id,
                            page_type="markdown", author=author,
                        )
                        results["success"].append({
                            "page_id": wiki_page.page_id,
                            "title": wiki_page.title,
                            "notion_id": page_id,
                        })
                    except Exception as e:
                        logger.error(f"Notion页面导入失败: {e}")
                        results["failed"].append({"error": str(e)})

            logger.info(f"Notion API导入完成: {len(results['success'])} 成功, {len(results['failed'])} 失败")
        except Exception as e:
            logger.error(f"Notion API连接失败: {e}")
            results["failed"].append({"error": str(e)})
        return results

    def _extract_notion_block_text(self, block: Dict, block_type: str) -> str:
        rich_texts = block.get(block_type, {}).get("rich_text", [])
        parts = []
        for rt in rich_texts:
            text = rt.get("plain_text", "")
            annotations = rt.get("annotations", {})
            if annotations.get("bold"):
                text = f"**{text}**"
            if annotations.get("italic"):
                text = f"*{text}*"
            if annotations.get("strikethrough"):
                text = f"~~{text}~~"
            if annotations.get("code"):
                text = f"`{text}`"
            href = rt.get("href")
            if href:
                text = f"[{text}]({href})"
            parts.append(text)
        return "".join(parts)

    async def import_obsidian(self, vault_dir: str, space_id: str = "default",
                              author: str = "system") -> List[WikiPage]:
        pages = []
        vault_path = Path(vault_dir)
        if not vault_path.exists():
            logger.error(f"Obsidian库目录不存在: {vault_dir}")
            return pages
        link_map = {}
        for md_file in vault_path.rglob("*.md"):
            rel_path = md_file.relative_to(vault_path)
            page_name = rel_path.stem
            link_map[page_name.lower()] = str(rel_path.with_suffix(""))
        for md_file in vault_path.rglob("*.md"):
            try:
                with open(md_file, "r", encoding="utf-8") as f:
                    content = f.read()
                content = self._convert_obsidian_links(content, link_map)
                content = re.sub(r'#(\w[\w\u4e00-\u9fff]*)', r'标签:\1', content)
                content = self._convert_obsidian_embeds(content)
                frontmatter = {}
                content, frontmatter = self._parse_obsidian_frontmatter(content)
                rel_path = md_file.relative_to(vault_path)
                title = str(rel_path.with_suffix(""))
                tags = frontmatter.get("tags", [])
                if isinstance(tags, str):
                    tags = [t.strip() for t in tags.split(",")]
                page = wiki_store.create_page(
                    title=title, content=content, space_id=space_id,
                    page_type="markdown", author=author, tags=tags,
                )
                pages.append(page)
            except Exception as e:
                logger.error(f"Obsidian文件导入失败 {md_file}: {e}")
        logger.info(f"Obsidian导入完成: {len(pages)} 个页面")
        return pages

    def _convert_obsidian_links(self, content: str, link_map: Dict[str, str]) -> str:
        def replace_wikilink(match):
            link_text = match.group(1)
            parts = link_text.split("|")
            target = parts[0].strip()
            alias = parts[1].strip() if len(parts) > 1 else target
            resolved = link_map.get(target.lower(), target)
            return f"[{alias}]({resolved})"
        content = re.sub(r'\[\[([^\]]+)\]\]', replace_wikilink, content)
        return content

    def _convert_obsidian_embeds(self, content: str) -> str:
        content = re.sub(r'!\[\[([^\]]+)\]\]', r'嵌入文件: \1', content)
        return content

    def _parse_obsidian_frontmatter(self, content: str) -> tuple:
        frontmatter = {}
        if content.startswith("---"):
            parts = content.split("---", 2)
            if len(parts) >= 3:
                try:
                    import yaml
                    fm_text = parts[1].strip()
                    frontmatter = yaml.safe_load(fm_text) or {}
                    content = parts[2].strip()
                except ImportError:
                    for line in parts[1].strip().split("\n"):
                        if ":" in line:
                            key, val = line.split(":", 1)
                            frontmatter[key.strip()] = val.strip()
                    content = parts[2].strip()
                except Exception:
                    pass
        return content, frontmatter

    async def _import_csv(self, file_path: str, space_id: str = "default",
                          author: str = "system") -> Optional[WikiPage]:
        try:
            import csv
            with open(file_path, "r", encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                rows = list(reader)
            if not rows:
                return None
            content_parts = []
            header = rows[0]
            content_parts.append("| " + " | ".join(header) + " |")
            content_parts.append("| " + " | ".join(["---"] * len(header)) + " |")
            for row in rows[1:]:
                content_parts.append("| " + " | ".join(row) + " |")
            content = "\n".join(content_parts)
            title = Path(file_path).stem
            return wiki_store.create_page(
                title=title, content=content, space_id=space_id,
                page_type="markdown", author=author,
            )
        except Exception as e:
            logger.error(f"CSV导入失败: {e}")
            return None

    async def batch_import_files(self, file_paths: List[str], space_id: str = "default",
                                 author: str = "system") -> Dict[str, Any]:
        results = {"success": [], "failed": []}
        for fp in file_paths:
            ext = Path(fp).suffix.lower()
            page = None
            if ext == ".pdf":
                page = await self.import_pdf(fp, space_id, author)
            elif ext in (".docx", ".doc"):
                page = await self.import_word(fp, space_id, author)
            elif ext in (".md", ".markdown"):
                page = await self.import_markdown(fp, space_id, author)
            elif ext in (".xlsx", ".xls"):
                page = await self.import_excel(fp, space_id, author)
            elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif"):
                page = await self.import_image_ocr(fp, space_id, author)
            elif ext == ".csv":
                page = await self._import_csv(fp, space_id, author)
            else:
                results["failed"].append({"file": fp, "error": f"不支持的格式: {ext}"})
                continue

            if page:
                results["success"].append({"file": fp, "page_id": page.page_id, "title": page.title})
            else:
                results["failed"].append({"file": fp, "error": "导入失败"})
        return results

    async def import_folder(self, folder_path: str, space_id: str = "default",
                            author: str = "system") -> Dict[str, Any]:
        folder = Path(folder_path)
        if not folder.exists():
            return {"success": [], "failed": [{"file": folder_path, "error": "目录不存在"}]}
        file_paths = []
        for ext in ["*.pdf", "*.docx", "*.doc", "*.md", "*.markdown", "*.xlsx", "*.xls",
                     "*.png", "*.jpg", "*.jpeg", "*.bmp", "*.tiff", "*.tif", "*.csv"]:
            file_paths.extend([str(f) for f in folder.rglob(ext)])
        if not file_paths:
            return {"success": [], "failed": []}
        return await self.batch_import_files(file_paths, space_id, author)


doc_import_service = DocumentImportService()
