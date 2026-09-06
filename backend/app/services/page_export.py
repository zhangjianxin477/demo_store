import os
import re
import io
import uuid
import json
import logging
from typing import Optional, Dict, Any
from datetime import datetime

from app.core.config import settings
from app.models.wiki import wiki_store
from app.services.page_embed import page_embed_service

logger = logging.getLogger(__name__)


class PageExportService:
    def __init__(self):
        self._export_dir = os.path.join(settings.DATA_DIR, "exports")
        os.makedirs(self._export_dir, exist_ok=True)

    def export_markdown(self, page_id: str, resolve_embeds: bool = True) -> Optional[Dict]:
        page = wiki_store.get_page(page_id)
        if not page:
            return None
        content = page.content
        if resolve_embeds:
            content = page_embed_service.resolve_embeds(content)
        content = f"# {page.title}\n\n{content}"
        filename = self._safe_filename(page.title) + ".md"
        filepath = os.path.join(self._export_dir, filename)
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(content)
        return {
            "success": True,
            "format": "markdown",
            "filename": filename,
            "filepath": filepath,
            "size": len(content.encode("utf-8")),
        }

    def export_html(self, page_id: str, resolve_embeds: bool = True,
                    include_styling: bool = True) -> Optional[Dict]:
        page = wiki_store.get_page(page_id)
        if not page:
            return None
        content = page.content
        if resolve_embeds:
            content = page_embed_service.resolve_embeds(content)

        html_body = self._markdown_to_html(content)
        if include_styling:
            html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{self._escape_html(page.title)}</title>
<style>
body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; max-width: 900px; margin: 0 auto; padding: 40px 20px; line-height: 1.8; color: #333; }}
h1 {{ border-bottom: 2px solid #6366f1; padding-bottom: 8px; color: #1a1a2e; }}
h2 {{ border-bottom: 1px solid #e5e7eb; padding-bottom: 6px; color: #374151; margin-top: 2em; }}
h3 {{ color: #4b5563; }}
table {{ border-collapse: collapse; width: 100%; margin: 16px 0; }}
th, td {{ border: 1px solid #d1d5db; padding: 8px 12px; text-align: left; }}
th {{ background: #f3f4f6; font-weight: 600; }}
blockquote {{ border-left: 4px solid #6366f1; margin: 16px 0; padding: 8px 16px; background: #f8f9fa; }}
code {{ background: #f1f5f9; padding: 2px 6px; border-radius: 3px; font-size: 0.9em; }}
pre {{ background: #1e293b; color: #e2e8f0; padding: 16px; border-radius: 8px; overflow-x: auto; }}
pre code {{ background: none; padding: 0; }}
mark {{ background: #fef08a; padding: 1px 4px; }}
img {{ max-width: 100%; }}
a {{ color: #6366f1; }}
</style>
</head>
<body>
<h1>{self._escape_html(page.title)}</h1>
{html_body}
<hr>
<p style="color:#9ca3af;font-size:12px">导出自 Knowledge Hub · {datetime.now().strftime('%Y-%m-%d %H:%M')}</p>
</body>
</html>"""
        else:
            html = f"<h1>{self._escape_html(page.title)}</h1>\n{html_body}"

        filename = self._safe_filename(page.title) + ".html"
        filepath = os.path.join(self._export_dir, filename)
        with open(filepath, "w", encoding="utf-8") as f:
            f.write(html)
        return {
            "success": True,
            "format": "html",
            "filename": filename,
            "filepath": filepath,
            "size": len(html.encode("utf-8")),
        }

    def export_pdf(self, page_id: str, resolve_embeds: bool = True) -> Optional[Dict]:
        page = wiki_store.get_page(page_id)
        if not page:
            return None

        try:
            from weasyprint import HTML
            html_result = self.export_html(page_id, resolve_embeds, include_styling=True)
            if not html_result:
                return None
            with open(html_result["filepath"], "r", encoding="utf-8") as f:
                html_content = f.read()
            filename = self._safe_filename(page.title) + ".pdf"
            filepath = os.path.join(self._export_dir, filename)
            HTML(string=html_content).write_pdf(filepath)
            return {
                "success": True,
                "format": "pdf",
                "filename": filename,
                "filepath": filepath,
                "size": os.path.getsize(filepath),
            }
        except ImportError:
            logger.warning("weasyprint未安装，PDF导出不可用，尝试使用其他方式")
            return self._export_pdf_alternative(page, resolve_embeds)
        except Exception as e:
            logger.error(f"PDF导出失败: {e}")
            return {"success": False, "error": str(e)}

    def _export_pdf_alternative(self, page, resolve_embeds: bool) -> Optional[Dict]:
        try:
            import subprocess
            html_result = self.export_html(page_id=page.page_id, resolve_embeds=resolve_embeds)
            if not html_result:
                return None
            filename = self._safe_filename(page.title) + ".pdf"
            filepath = os.path.join(self._export_dir, filename)
            try:
                subprocess.run(
                    ["wkhtmltopdf", "--encoding", "utf-8", html_result["filepath"], filepath],
                    capture_output=True, timeout=30,
                )
                if os.path.exists(filepath):
                    return {
                        "success": True,
                        "format": "pdf",
                        "filename": filename,
                        "filepath": filepath,
                        "size": os.path.getsize(filepath),
                    }
            except Exception:
                pass
            return {
                "success": False,
                "error": "PDF导出需要安装weasyprint或wkhtmltopdf",
                "fallback": html_result,
            }
        except Exception as e:
            return {"success": False, "error": str(e)}

    def export_word(self, page_id: str, resolve_embeds: bool = True) -> Optional[Dict]:
        page = wiki_store.get_page(page_id)
        if not page:
            return None
        content = page.content
        if resolve_embeds:
            content = page_embed_service.resolve_embeds(content)

        try:
            from docx import Document
            from docx.shared import Pt, Inches
            doc = Document()
            doc.add_heading(page.title, level=0)
            for line in content.split("\n"):
                stripped = line.strip()
                if not stripped:
                    continue
                if stripped.startswith("# "):
                    doc.add_heading(stripped[2:], level=1)
                elif stripped.startswith("## "):
                    doc.add_heading(stripped[3:], level=2)
                elif stripped.startswith("### "):
                    doc.add_heading(stripped[4:], level=3)
                elif stripped.startswith("- ") or stripped.startswith("* "):
                    doc.add_paragraph(stripped[2:], style="List Bullet")
                elif stripped.startswith("> "):
                    p = doc.add_paragraph(stripped[2:])
                    p.style = doc.styles.get_by_name("Intense Quote") if "Intense Quote" in [s.name for s in doc.styles] else p.style
                elif stripped.startswith("| ") and "|" in stripped[1:]:
                    self._add_table_to_doc(doc, stripped, content)
                    break
                else:
                    doc.add_paragraph(stripped)

            filename = self._safe_filename(page.title) + ".docx"
            filepath = os.path.join(self._export_dir, filename)
            doc.save(filepath)
            return {
                "success": True,
                "format": "word",
                "filename": filename,
                "filepath": filepath,
                "size": os.path.getsize(filepath),
            }
        except ImportError:
            logger.warning("python-docx未安装，Word导出不可用")
            return {"success": False, "error": "Word导出需要安装python-docx"}
        except Exception as e:
            logger.error(f"Word导出失败: {e}")
            return {"success": False, "error": str(e)}

    def _add_table_to_doc(self, doc, first_line, content):
        pass

    def _markdown_to_html(self, md_text: str) -> str:
        try:
            import markdown
            return markdown.markdown(md_text, extensions=["tables", "fenced_code", "toc", "codehilite"])
        except ImportError:
            pass

        html = md_text
        html = re.sub(r'^### (.+)$', r'<h3>\1</h3>', html, flags=re.MULTILINE)
        html = re.sub(r'^## (.+)$', r'<h2>\1</h2>', html, flags=re.MULTILINE)
        html = re.sub(r'^# (.+)$', r'<h1>\1</h1>', html, flags=re.MULTILINE)
        html = re.sub(r'\*\*(.+?)\*\*', r'<strong>\1</strong>', html)
        html = re.sub(r'\*(.+?)\*', r'<em>\1</em>', html)
        html = re.sub(r'`(.+?)`', r'<code>\1</code>', html)
        html = re.sub(r'^> (.+)$', r'<blockquote>\1</blockquote>', html, flags=re.MULTILINE)
        html = re.sub(r'^- (.+)$', r'<li>\1</li>', html, flags=re.MULTILINE)
        html = re.sub(r'^\d+\. (.+)$', r'<li>\1</li>', html, flags=re.MULTILINE)
        html = re.sub(r'\n{2,}', '</p><p>', html)
        html = f'<p>{html}</p>'
        return html

    def _safe_filename(self, title: str) -> str:
        safe = re.sub(r'[\\/:*?"<>|]', '_', title)
        return safe[:100] if safe else "export"

    def _escape_html(self, text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


page_export_service = PageExportService()
