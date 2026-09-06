import os
import re
import logging
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass, field
from enum import Enum

logger = logging.getLogger(__name__)


class LayoutBlockType(str, Enum):
    TITLE = "title"
    HEADING_H1 = "heading_h1"
    HEADING_H2 = "heading_h2"
    HEADING_H3 = "heading_h3"
    HEADING_H4 = "heading_h4"
    PARAGRAPH = "paragraph"
    TABLE = "table"
    FIGURE = "figure"
    IMAGE = "image"
    CAPTION = "caption"
    FOOTNOTE = "footnote"
    HEADER = "header"
    FOOTER = "footer"
    PAGE_NUMBER = "page_number"
    CODE_BLOCK = "code_block"
    LIST_ITEM = "list_item"
    QUOTE = "quote"
    ABSTRACT = "abstract"
    REFERENCES = "references"
    FORMULA = "formula"
    UNKNOWN = "unknown"


@dataclass
class LayoutBlock:
    block_type: LayoutBlockType
    content: str
    page: int = 0
    bbox: Optional[Tuple[float, float, float, float]] = None
    confidence: float = 1.0
    level: int = 0
    metadata: Dict[str, Any] = field(default_factory=dict)
    children: List['LayoutBlock'] = field(default_factory=list)


@dataclass
class LayoutResult:
    blocks: List[LayoutBlock] = field(default_factory=list)
    page_count: int = 0
    images: List[Dict[str, Any]] = field(default_factory=list)
    tables: List[Dict[str, Any]] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)


class LayoutAnalyzer:
    _HEADING_PATTERNS = [
        (re.compile(r'^#{1}\s+'), LayoutBlockType.HEADING_H1, 1),
        (re.compile(r'^#{2}\s+'), LayoutBlockType.HEADING_H2, 2),
        (re.compile(r'^#{3}\s+'), LayoutBlockType.HEADING_H3, 3),
        (re.compile(r'^#{4,}\s+'), LayoutBlockType.HEADING_H4, 4),
        (re.compile(r'^(\d+(?:\.\d+)*)\s+[、.．]\s*\S'), LayoutBlockType.HEADING_H2, 2),
        (re.compile(r'^(\d+\.\d+)\s+\S'), LayoutBlockType.HEADING_H3, 3),
        (re.compile(r'^(\d+\.\d+\.\d+)\s+\S'), LayoutBlockType.HEADING_H4, 4),
        (re.compile(r'^第[一二三四五六七八九十百千\d]+[章回编部篇]'), LayoutBlockType.HEADING_H1, 1),
        (re.compile(r'^第[一二三四五六七八九十百千\d]+[节条款]'), LayoutBlockType.HEADING_H2, 2),
        (re.compile(r'^(?:一|二|三|四|五|六|七|八|九|十)+[、．.]\s'), LayoutBlockType.HEADING_H2, 2),
        (re.compile(r'^[（(]\s*[一二三四五六七八九十\d]+\s*[）)]\s'), LayoutBlockType.HEADING_H3, 3),
    ]

    _TITLE_KEYWORDS = [
        "摘要", "Abstract", "ABSTRACT",
        "关键词", "Keywords", "KEYWORDS",
        "目录", "Contents", "引言", "前言", "绪论",
        "结论", "结语", "参考文献", "References",
        "致谢", "Acknowledgement", "附录", "Appendix",
    ]

    _TABLE_PATTERN = re.compile(r'^\|.+\|$', re.MULTILINE)
    _TABLE_SEP_PATTERN = re.compile(r'^\|[\s\-:]+\|', re.MULTILINE)
    _CODE_FENCE_START = re.compile(r'^```')
    _CODE_FENCE_END = re.compile(r'^```$')
    _IMAGE_PATTERN = re.compile(r'!\[.*?\]\(.*?\)')
    _LIST_PATTERN = re.compile(r'^[\s]*[-*+•]\s|^\d+[\.、）)]\s', re.MULTILINE)
    _QUOTE_PATTERN = re.compile(r'^>\s', re.MULTILINE)
    _FORMULA_PATTERN = re.compile(r'^\$\$|^\s*\\begin\{', re.MULTILINE)
    _FOOTNOTE_PATTERN = re.compile(r'^\[\^[^\]]+\]:|^\[\d+\]', re.MULTILINE)
    _PAGE_NUM_PATTERN = re.compile(r'^\s*\d+\s*$')

    _yolo_model = None
    _yolo_loaded = False

    @classmethod
    def _try_load_yolo(cls):
        if cls._yolo_loaded:
            return cls._yolo_model is not None
        cls._yolo_loaded = True
        try:
            from ultralytics import YOLO
            model_path = os.environ.get("LAYOUT_YOLO_MODEL", "doclayout_yolo_docstructbench_imgsz640.pt")
            if os.path.exists(model_path):
                cls._yolo_model = YOLO(model_path)
                logger.info(f"DocLayout-YOLO 模型加载成功: {model_path}")
                return True
            logger.info(f"YOLO 模型文件不存在: {model_path}，使用规则版面分析")
            return False
        except ImportError:
            logger.info("ultralytics 未安装，跳过 YOLO 版面分析")
            return False
        except Exception as e:
            logger.warning(f"YOLO 模型加载失败: {e}，使用规则版面分析")
            return False

    def analyze_pdf(self, file_path: str) -> LayoutResult:
        if self._try_load_yolo():
            try:
                return self._analyze_pdf_with_yolo(file_path)
            except Exception as e:
                logger.warning(f"YOLO 版面分析失败: {e}，回退到规则分析")

        return self._analyze_pdf_rule_based(file_path)

    def analyze_text(self, text: str, filename: str = "") -> LayoutResult:
        return self._analyze_text_rule_based(text, filename)

    def analyze_content(self, content: bytes, filename: str) -> LayoutResult:
        ext = os.path.splitext(filename)[1].lower()
        if ext == '.pdf':
            import tempfile
            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            try:
                return self.analyze_pdf(tmp_path)
            finally:
                os.unlink(tmp_path)
        elif ext in ('.docx', '.doc'):
            return self._analyze_docx_content(content, filename)
        else:
            try:
                text = content.decode('utf-8')
            except UnicodeDecodeError:
                try:
                    text = content.decode('gbk')
                except UnicodeDecodeError:
                    text = content.decode('utf-8', errors='replace')
            return self.analyze_text(text, filename)

    def _analyze_pdf_with_yolo(self, file_path: str) -> LayoutResult:
        import fitz
        doc = fitz.open(file_path)
        result = LayoutResult(page_count=len(doc))

        for page_idx in range(len(doc)):
            page = doc[page_idx]
            pix = page.get_pixmap(dpi=150)
            import numpy as np
            img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)

            yolo_results = self._yolo_model.predict(img, imgsz=640, conf=0.25, verbose=False)

            for r in yolo_results:
                for box in r.boxes:
                    x1, y1, x2, y2 = box.xyxy[0].tolist()
                    cls_id = int(box.cls[0])
                    conf = float(box.conf[0])
                    class_name = r.names.get(cls_id, "unknown")

                    block_type = self._yolo_class_to_block_type(class_name)
                    rect = fitz.Rect(x1 * page.rect.width / pix.width,
                                     y1 * page.rect.height / pix.height,
                                     x2 * page.rect.width / pix.width,
                                     y2 * page.rect.height / pix.height)
                    text = page.get_text("text", clip=rect).strip()

                    block = LayoutBlock(
                        block_type=block_type,
                        content=text,
                        page=page_idx + 1,
                        bbox=(x1, y1, x2, y2),
                        confidence=conf,
                        level=self._block_type_to_level(block_type),
                    )
                    result.blocks.append(block)

                    if block_type == LayoutBlockType.TABLE:
                        result.tables.append({
                            "page": page_idx + 1,
                            "bbox": (x1, y1, x2, y2),
                            "content": text,
                        })
                    elif block_type in (LayoutBlockType.FIGURE, LayoutBlockType.IMAGE):
                        result.images.append({
                            "page": page_idx + 1,
                            "bbox": (x1, y1, x2, y2),
                        })

        doc.close()
        return result

    def _analyze_pdf_rule_based(self, file_path: str) -> LayoutResult:
        try:
            import fitz
            doc = fitz.open(file_path)
        except ImportError:
            try:
                import pdfplumber
                return self._analyze_pdf_with_pdfplumber(file_path)
            except ImportError:
                logger.warning("fitz 和 pdfplumber 均不可用，使用纯文本分析")
                with open(file_path, 'rb') as f:
                    text = f.read().decode('utf-8', errors='replace')
                return self.analyze_text(text, os.path.basename(file_path))

        result = LayoutResult(page_count=len(doc))
        page_layouts = []

        for page_idx in range(len(doc)):
            page = doc[page_idx]
            blocks = page.get_text("blocks")

            # 检测分栏布局：分析文本块的x坐标分布
            column_info = self._detect_pdf_columns(blocks, page.rect.width)
            page_layouts.append({
                "page": page_idx + 1,
                "is_multi_column": bool(column_info.get("is_multi_column")),
                "num_columns": column_info.get("num_columns", 1),
            })

            page_table_bboxes = []
            try:
                tables = page.find_tables()
            except Exception:
                tables = []

            for table in tables:
                table_dict = table.extract()
                if table_dict:
                    table_md = self._table_to_markdown(table_dict)
                    page_table_bboxes.append(table.bbox)
                    result.tables.append({
                        "page": page_idx + 1,
                        "bbox": table.bbox,
                        "content": table_md,
                        "rows": len(table_dict),
                        "cols": len(table_dict[0]) if table_dict else 0,
                    })
                    result.blocks.append(LayoutBlock(
                        block_type=LayoutBlockType.TABLE,
                        content=table_md,
                        page=page_idx + 1,
                        bbox=table.bbox,
                        metadata={
                            "rows": len(table_dict),
                            "cols": len(table_dict[0]) if table_dict else 0,
                            "column": self._get_column_index(table.bbox[0], table.bbox[2], column_info),
                            "source": "pdf_table_detector",
                        },
                    ))

            for block in blocks:
                if len(block) < 5:
                    continue
                x0, y0, x1, y1, block_text, block_no, block_type = block[:7]
                block_bbox = (x0, y0, x1, y1)
                if any(self._bbox_overlap_ratio(block_bbox, table_bbox) > 0.55 for table_bbox in page_table_bboxes):
                    continue

                if block_type == 1:
                    result.images.append({
                        "page": page_idx + 1,
                        "bbox": (x0, y0, x1, y1),
                    })
                    result.blocks.append(LayoutBlock(
                        block_type=LayoutBlockType.IMAGE,
                        content=f"[图片: 第{page_idx + 1}页]",
                        page=page_idx + 1,
                        bbox=block_bbox,
                        metadata={
                            "column": self._get_column_index(x0, x1, column_info),
                            "source": "pdf_image_block",
                        },
                    ))
                    continue

                block_text = block_text.strip()
                if not block_text:
                    continue

                # 记录该块所属的栏
                col_idx = self._get_column_index(x0, x1, column_info)

                lines = block_text.split('\n')
                for line in lines:
                    line = line.strip()
                    if not line:
                        continue

                    block_type_detected = self._classify_line(line)
                    level = self._block_type_to_level(block_type_detected)

                    result.blocks.append(LayoutBlock(
                        block_type=block_type_detected,
                        content=line,
                        page=page_idx + 1,
                        bbox=block_bbox,
                        level=level,
                        metadata={"column": col_idx},
                    ))

        doc.close()

        # 对分栏页面进行内容重排序：确保同一栏的内容连续，避免跨栏内容被错误分割
        result.blocks = self._reorder_column_blocks(result.blocks)
        result.metadata["layout_pages"] = page_layouts
        result.metadata["has_multi_column"] = any(p["is_multi_column"] for p in page_layouts)
        result.metadata["tables_count"] = len(result.tables)
        result.metadata["images_count"] = len(result.images)

        return result

    @staticmethod
    def _bbox_overlap_ratio(a: Tuple[float, float, float, float], b: Tuple[float, float, float, float]) -> float:
        ax0, ay0, ax1, ay1 = a
        bx0, by0, bx1, by1 = b
        ix0 = max(ax0, bx0)
        iy0 = max(ay0, by0)
        ix1 = min(ax1, bx1)
        iy1 = min(ay1, by1)
        if ix1 <= ix0 or iy1 <= iy0:
            return 0.0
        intersection = (ix1 - ix0) * (iy1 - iy0)
        area = max((ax1 - ax0) * (ay1 - ay0), 1.0)
        return intersection / area

    @staticmethod
    def _detect_pdf_columns(blocks: list, page_width: float) -> Dict[str, Any]:
        """检测PDF页面是否为分栏布局，返回分栏信息。
        通过分析文本块的x坐标分布，判断是否存在明显的分栏间隙。"""
        if page_width < 100:
            return {"is_multi_column": False, "boundaries": []}

        # 收集文本块的x坐标
        text_blocks_x = []
        for block in blocks:
            if len(block) < 5:
                continue
            x0, y0, x1, y1, block_text, block_no, block_type = block[:7]
            if block_type == 1:  # 图片块跳过
                continue
            if not block_text or not block_text.strip():
                continue
            text_blocks_x.append((x0, x1))

        if len(text_blocks_x) < 4:
            return {"is_multi_column": False, "boundaries": []}

        # 统计x坐标分布，检测中间间隙
        num_bins = 40
        bin_width = page_width / num_bins
        bins = [0] * num_bins
        for x0, x1 in text_blocks_x:
            mid = (x0 + x1) / 2
            bin_idx = int(mid / bin_width)
            if 0 <= bin_idx < num_bins:
                bins[bin_idx] += 1

        max_bin = max(bins) if bins else 0
        if max_bin == 0:
            return {"is_multi_column": False, "boundaries": []}

        # 寻找中间区域的低密度间隙
        threshold = max_bin * 0.15
        center_start = num_bins // 4
        center_end = num_bins * 3 // 4
        gap_bins = [i for i in range(center_start, center_end) if bins[i] < threshold]

        if not gap_bins:
            return {"is_multi_column": False, "boundaries": []}

        # 将连续的间隙bin合并为间隙区域
        groups = []
        current_group = [gap_bins[0]]
        for i in range(1, len(gap_bins)):
            if gap_bins[i] - gap_bins[i - 1] <= 2:
                current_group.append(gap_bins[i])
            else:
                groups.append(current_group)
                current_group = [gap_bins[i]]
        groups.append(current_group)

        # 找到最宽的间隙区域作为分栏边界
        best_group = max(groups, key=len) if groups else None
        if best_group is None or len(best_group) < 2:
            return {"is_multi_column": False, "boundaries": []}

        # 计算分栏边界坐标
        boundary = page_width * (best_group[0] + best_group[-1] + 1) / 2 / num_bins

        # 验证：左右两侧都应有足够的文本块
        left_count = sum(1 for x0, x1 in text_blocks_x if x1 < boundary)
        right_count = sum(1 for x0, x1 in text_blocks_x if x0 >= boundary)

        if left_count < 3 or right_count < 3:
            return {"is_multi_column": False, "boundaries": []}

        return {
            "is_multi_column": True,
            "boundaries": [0, boundary, page_width],
            "num_columns": 2,
        }

    @staticmethod
    def _get_column_index(x0: float, x1: float, column_info: Dict[str, Any]) -> int:
        """根据文本块的x坐标和分栏信息，判断该块属于哪一栏"""
        if not column_info.get("is_multi_column"):
            return 0
        boundaries = column_info.get("boundaries", [])
        mid_x = (x0 + x1) / 2
        for i in range(len(boundaries) - 1):
            if boundaries[i] <= mid_x < boundaries[i + 1]:
                return i
        return len(boundaries) - 2 if len(boundaries) >= 2 else 0

    @staticmethod
    def _reorder_column_blocks(blocks: List['LayoutBlock']) -> List['LayoutBlock']:
        """对分栏页面的布局块进行重排序，确保同一栏的内容连续排列。
        避免跨栏内容在分块时被错误分割。
        排序规则：按页码 → 按栏号 → 按y坐标（从上到下）"""
        if not blocks:
            return blocks

        # 检查是否有分栏信息
        has_columns = any(b.metadata.get("column", 0) > 0 for b in blocks)
        if not has_columns:
            return blocks

        # 按页码分组，每组内按栏号和y坐标排序
        from itertools import groupby
        sorted_blocks = sorted(blocks, key=lambda b: (
            b.page,
            b.metadata.get("column", 0),
            b.bbox[1] if b.bbox else 0,  # y0坐标
        ))
        return sorted_blocks

    def _analyze_pdf_with_pdfplumber(self, file_path: str) -> LayoutResult:
        import pdfplumber
        result = LayoutResult()
        with pdfplumber.open(file_path) as pdf:
            result.page_count = len(pdf.pages)
            for page_idx, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                for line in text.split('\n'):
                    line = line.strip()
                    if not line:
                        continue
                    block_type = self._classify_line(line)
                    result.blocks.append(LayoutBlock(
                        block_type=block_type,
                        content=line,
                        page=page_idx + 1,
                        level=self._block_type_to_level(block_type),
                    ))

                tables = page.extract_tables()
                for table in tables:
                    if table:
                        table_md = self._table_to_markdown(table)
                        result.tables.append({
                            "page": page_idx + 1,
                            "content": table_md,
                        })
                        result.blocks.append(LayoutBlock(
                            block_type=LayoutBlockType.TABLE,
                            content=table_md,
                            page=page_idx + 1,
                        ))

        return result

    def _analyze_text_rule_based(self, text: str, filename: str = "") -> LayoutResult:
        result = LayoutResult(page_count=1)
        if not text or not text.strip():
            return result

        lines = text.split('\n')
        in_code_block = False
        code_buffer = []
        code_lang = ""
        table_buffer = []
        in_table = False

        i = 0
        while i < len(lines):
            line = lines[i]
            stripped = line.strip()

            if self._CODE_FENCE_START.match(stripped) and not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_buffer = [line]
                i += 1
                continue

            if in_code_block:
                code_buffer.append(line)
                if self._CODE_FENCE_END.match(stripped):
                    in_code_block = False
                    result.blocks.append(LayoutBlock(
                        block_type=LayoutBlockType.CODE_BLOCK,
                        content='\n'.join(code_buffer),
                        metadata={"lang": code_lang},
                    ))
                    code_buffer = []
                i += 1
                continue

            if self._is_table_row(stripped):
                if not in_table:
                    in_table = True
                    table_buffer = []
                table_buffer.append(stripped)
                i += 1
                continue
            elif in_table:
                in_table = False
                table_md = '\n'.join(table_buffer)
                result.tables.append({"page": 1, "content": table_md})
                result.blocks.append(LayoutBlock(
                    block_type=LayoutBlockType.TABLE,
                    content=table_md,
                    page=1,
                ))
                table_buffer = []

            if not stripped:
                i += 1
                continue

            if self._IMAGE_PATTERN.search(stripped):
                result.images.append({"page": 1, "content": stripped})
                result.blocks.append(LayoutBlock(
                    block_type=LayoutBlockType.IMAGE,
                    content=stripped,
                    page=1,
                ))
                i += 1
                continue

            if self._FORMULA_PATTERN.match(stripped):
                formula_lines = [line]
                j = i + 1
                while j < len(lines):
                    next_stripped = lines[j].strip()
                    formula_lines.append(lines[j])
                    if next_stripped.endswith('$$') or re.match(r'^\s*\\end\{', next_stripped):
                        j += 1
                        break
                    j += 1
                result.blocks.append(LayoutBlock(
                    block_type=LayoutBlockType.FORMULA,
                    content='\n'.join(formula_lines),
                    page=1,
                ))
                i = j
                continue

            if self._FOOTNOTE_PATTERN.match(stripped):
                result.blocks.append(LayoutBlock(
                    block_type=LayoutBlockType.FOOTNOTE,
                    content=stripped,
                    page=1,
                ))
                i += 1
                continue

            if self._QUOTE_PATTERN.match(stripped):
                result.blocks.append(LayoutBlock(
                    block_type=LayoutBlockType.QUOTE,
                    content=stripped,
                    page=1,
                ))
                i += 1
                continue

            block_type = self._classify_line(stripped)
            level = self._block_type_to_level(block_type)

            if block_type in (LayoutBlockType.HEADING_H1, LayoutBlockType.HEADING_H2,
                              LayoutBlockType.HEADING_H3, LayoutBlockType.HEADING_H4,
                              LayoutBlockType.TITLE):
                result.blocks.append(LayoutBlock(
                    block_type=block_type,
                    content=stripped,
                    page=1,
                    level=level,
                ))
            else:
                para_lines = [stripped]
                j = i + 1
                while j < len(lines):
                    next_stripped = lines[j].strip()
                    if not next_stripped:
                        break
                    if self._classify_line(next_stripped) in (LayoutBlockType.HEADING_H1,
                                                               LayoutBlockType.HEADING_H2,
                                                               LayoutBlockType.HEADING_H3,
                                                               LayoutBlockType.HEADING_H4,
                                                               LayoutBlockType.TITLE):
                        break
                    if self._is_table_row(next_stripped):
                        break
                    if self._CODE_FENCE_START.match(next_stripped):
                        break
                    if self._IMAGE_PATTERN.search(next_stripped):
                        break
                    para_lines.append(next_stripped)
                    j += 1

                result.blocks.append(LayoutBlock(
                    block_type=LayoutBlockType.PARAGRAPH,
                    content='\n'.join(para_lines),
                    page=1,
                ))
                i = j - 1

            i += 1

        if in_table and table_buffer:
            table_md = '\n'.join(table_buffer)
            result.tables.append({"page": 1, "content": table_md})
            result.blocks.append(LayoutBlock(
                block_type=LayoutBlockType.TABLE,
                content=table_md,
                page=1,
            ))

        return result

    def _analyze_docx_content(self, content: bytes, filename: str) -> LayoutResult:
        result = LayoutResult(page_count=1)
        try:
            from docx import Document
            import io
            doc = Document(io.BytesIO(content))

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                style_name = (para.style.name or "").lower() if para.style else ""
                block_type = LayoutBlockType.PARAGRAPH
                level = 0

                if "heading 1" in style_name or "title" in style_name:
                    block_type = LayoutBlockType.HEADING_H1
                    level = 1
                elif "heading 2" in style_name:
                    block_type = LayoutBlockType.HEADING_H2
                    level = 2
                elif "heading 3" in style_name:
                    block_type = LayoutBlockType.HEADING_H3
                    level = 3
                elif "heading 4" in style_name or "heading 5" in style_name:
                    block_type = LayoutBlockType.HEADING_H4
                    level = 4
                elif "caption" in style_name:
                    block_type = LayoutBlockType.CAPTION
                elif "quote" in style_name or "intense quote" in style_name:
                    block_type = LayoutBlockType.QUOTE
                elif "footnote" in style_name or "endnote" in style_name:
                    block_type = LayoutBlockType.FOOTNOTE
                else:
                    block_type = self._classify_line(text)
                    level = self._block_type_to_level(block_type)

                result.blocks.append(LayoutBlock(
                    block_type=block_type,
                    content=text,
                    page=1,
                    level=level,
                    metadata={"style": style_name},
                ))

            for table in doc.tables:
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(cells)
                if rows:
                    table_md = self._table_to_markdown(rows)
                    result.tables.append({"page": 1, "content": table_md, "rows": len(rows)})
                    result.blocks.append(LayoutBlock(
                        block_type=LayoutBlockType.TABLE,
                        content=table_md,
                        page=1,
                        metadata={"rows": len(rows), "cols": len(rows[0]) if rows else 0},
                    ))

            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    result.images.append({"page": 1, "rel_id": rel.rId})

        except Exception as e:
            logger.error(f"DOCX 版面分析失败: {e}")
            try:
                text = content.decode('utf-8', errors='replace')
                return self.analyze_text(text, filename)
            except Exception:
                pass

        return result

    def _classify_line(self, line: str) -> LayoutBlockType:
        if not line.strip():
            return LayoutBlockType.PARAGRAPH

        for pattern, block_type, _ in self._HEADING_PATTERNS:
            if pattern.match(line):
                return block_type

        if len(line.strip()) < 40:
            upper = line.strip().upper()
            for kw in self._TITLE_KEYWORDS:
                if upper.startswith(kw.upper()):
                    return LayoutBlockType.TITLE

        if self._PAGE_NUM_PATTERN.match(line) and len(line.strip()) <= 4:
            return LayoutBlockType.PAGE_NUMBER

        return LayoutBlockType.PARAGRAPH

    @staticmethod
    def _is_table_row(line: str) -> bool:
        return bool(LayoutAnalyzer._TABLE_PATTERN.match(line) or
                    LayoutAnalyzer._TABLE_SEP_PATTERN.match(line))

    @staticmethod
    def _block_type_to_level(block_type: LayoutBlockType) -> int:
        level_map = {
            LayoutBlockType.TITLE: 0,
            LayoutBlockType.HEADING_H1: 1,
            LayoutBlockType.HEADING_H2: 2,
            LayoutBlockType.HEADING_H3: 3,
            LayoutBlockType.HEADING_H4: 4,
        }
        return level_map.get(block_type, 0)

    @staticmethod
    def _yolo_class_to_block_type(class_name: str) -> LayoutBlockType:
        mapping = {
            "title": LayoutBlockType.TITLE,
            "text": LayoutBlockType.PARAGRAPH,
            "abandon": LayoutBlockType.UNKNOWN,
            "figure": LayoutBlockType.FIGURE,
            "figure_caption": LayoutBlockType.CAPTION,
            "table": LayoutBlockType.TABLE,
            "table_caption": LayoutBlockType.CAPTION,
            "table_footnote": LayoutBlockType.FOOTNOTE,
            "isolate_formula": LayoutBlockType.FORMULA,
            "formula_caption": LayoutBlockType.CAPTION,
            "header": LayoutBlockType.HEADER,
            "footer": LayoutBlockType.FOOTER,
            "reference": LayoutBlockType.REFERENCES,
            "page_number": LayoutBlockType.PAGE_NUMBER,
            "list": LayoutBlockType.LIST_ITEM,
            "code": LayoutBlockType.CODE_BLOCK,
            "caption": LayoutBlockType.CAPTION,
            "footnote": LayoutBlockType.FOOTNOTE,
            "image": LayoutBlockType.IMAGE,
        }
        return mapping.get(class_name.lower(), LayoutBlockType.UNKNOWN)

    @staticmethod
    def _table_to_markdown(table_data: List[List[str]]) -> str:
        if not table_data:
            return ""
        max_cols = max(len(row) for row in table_data)
        normalized = []
        for row in table_data:
            padded = [str(cell).strip() if cell else "" for cell in row]
            while len(padded) < max_cols:
                padded.append("")
            normalized.append(padded)

        lines = []
        header = normalized[0]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join(["---"] * max_cols) + " |")
        for row in normalized[1:]:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)


layout_analyzer = LayoutAnalyzer()
