import os
import io
import uuid
import hashlib
import logging
from typing import List, Dict, Any, Optional

from app.core.config import settings

logger = logging.getLogger(__name__)


class ImageExtractor:
    def __init__(self):
        self._image_dir = os.path.join(settings.DATA_DIR, "images")
        os.makedirs(self._image_dir, exist_ok=True)

    def extract_from_pdf(self, file_path: str) -> List[Dict[str, Any]]:
        images = []
        try:
            import fitz
            doc = fitz.open(file_path)
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                image_list = page.get_images(full=True)
                for img_idx, img_info in enumerate(image_list):
                    xref = img_info[0]
                    try:
                        base_image = doc.extract_image(xref)
                        if not base_image:
                            continue
                        image_bytes = base_image["image"]
                        image_ext = base_image.get("ext", "png")
                        width = base_image.get("width", 0)
                        height = base_image.get("height", 0)

                        if width < 50 or height < 50:
                            continue

                        saved_path = self._save_image(image_bytes, image_ext, page_idx, img_idx)
                        if saved_path:
                            images.append({
                                "page": page_idx + 1,
                                "index": img_idx,
                                "path": saved_path,
                                "width": width,
                                "height": height,
                                "format": image_ext,
                                "size_bytes": len(image_bytes),
                                "caption": f"图{len(images) + 1}: 第{page_idx + 1}页图片{img_idx + 1}",
                            })
                    except Exception as e:
                        logger.warning(f"PDF 图片提取失败 (page={page_idx}, img={img_idx}): {e}")
            doc.close()
        except ImportError:
            logger.info("fitz 不可用，尝试 pdfplumber 提取图片")
            images = self._extract_from_pdf_pdfplumber(file_path)
        except Exception as e:
            logger.error(f"PDF 图片提取失败: {e}")
        return images

    def _extract_from_pdf_pdfplumber(self, file_path: str) -> List[Dict[str, Any]]:
        images = []
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    if hasattr(page, 'images') and page.images:
                        for img_idx, img in enumerate(page.images):
                            images.append({
                                "page": page_idx + 1,
                                "index": img_idx,
                                "bbox": (img.get('x0', 0), img.get('top', 0),
                                         img.get('x1', 0), img.get('bottom', 0)),
                                "caption": f"图{len(images) + 1}: 第{page_idx + 1}页图片{img_idx + 1}",
                            })
        except Exception as e:
            logger.error(f"pdfplumber 图片提取失败: {e}")
        return images

    def extract_from_docx(self, content: bytes) -> List[Dict[str, Any]]:
        images = []
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))

            for rel_id, rel in doc.part.rels.items():
                if "image" not in rel.reltype:
                    continue
                try:
                    image_part = rel.target_part
                    image_bytes = image_part.blob
                    content_type = image_part.content_type
                    ext = self._content_type_to_ext(content_type)
                    img_idx = len(images)

                    saved_path = self._save_image(image_bytes, ext, 0, img_idx)
                    if saved_path:
                        images.append({
                            "page": 1,
                            "index": img_idx,
                            "path": saved_path,
                            "format": ext,
                            "size_bytes": len(image_bytes),
                            "rel_id": rel_id,
                            "caption": f"图{img_idx + 1}",
                        })
                except Exception as e:
                    logger.warning(f"DOCX 图片提取失败 ({rel_id}): {e}")

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if run._element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
                                    pass

        except Exception as e:
            logger.error(f"DOCX 图片提取失败: {e}")
        return images

    def extract_from_pptx(self, content: bytes) -> List[Dict[str, Any]]:
        images = []
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(content))
            for slide_idx, slide in enumerate(prs.slides):
                for shape in slide.shapes:
                    if shape.shape_type == 13:
                        try:
                            image = shape.image
                            image_bytes = image.blob
                            content_type = image.content_type
                            ext = self._content_type_to_ext(content_type)
                            img_idx = len(images)

                            saved_path = self._save_image(image_bytes, ext, slide_idx, img_idx)
                            if saved_path:
                                images.append({
                                    "page": slide_idx + 1,
                                    "index": img_idx,
                                    "path": saved_path,
                                    "format": ext,
                                    "size_bytes": len(image_bytes),
                                    "width": shape.width,
                                    "height": shape.height,
                                    "caption": f"图{img_idx + 1}: 幻灯片{slide_idx + 1}",
                                })
                        except Exception as e:
                            logger.warning(f"PPTX 图片提取失败 (slide={slide_idx}): {e}")
        except Exception as e:
            logger.error(f"PPTX 图片提取失败: {e}")
        return images

    def extract_from_html(self, content: bytes) -> List[Dict[str, Any]]:
        images = []
        try:
            from bs4 import BeautifulSoup
            for encoding in ['utf-8', 'gbk', 'latin-1']:
                try:
                    text = content.decode(encoding)
                    break
                except (UnicodeDecodeError, UnicodeError):
                    continue
            else:
                text = content.decode('utf-8', errors='replace')

            soup = BeautifulSoup(text, 'html.parser')
            for img_idx, img_tag in enumerate(soup.find_all('img')):
                src = img_tag.get('src', '')
                alt = img_tag.get('alt', f'图片{img_idx + 1}')
                if src:
                    images.append({
                        "index": img_idx,
                        "src": src,
                        "alt": alt,
                        "caption": f"图{img_idx + 1}: {alt}",
                    })
        except Exception as e:
            logger.error(f"HTML 图片提取失败: {e}")
        return images

    def extract_from_markdown(self, text: str) -> List[Dict[str, Any]]:
        import re
        images = []
        pattern = re.compile(r'!\[([^\]]*)\]\(([^)]+)\)')
        for match in pattern.finditer(text):
            alt = match.group(1)
            src = match.group(2)
            img_idx = len(images)
            images.append({
                "index": img_idx,
                "src": src,
                "alt": alt,
                "caption": f"图{img_idx + 1}: {alt}" if alt else f"图{img_idx + 1}",
            })
        return images

    def _save_image(self, image_bytes: bytes, ext: str, page: int, index: int) -> Optional[str]:
        try:
            content_hash = hashlib.md5(image_bytes).hexdigest()[:12]
            filename = f"img_p{page}_{index}_{content_hash}.{ext}"
            filepath = os.path.join(self._image_dir, filename)

            if os.path.exists(filepath):
                return filepath

            with open(filepath, 'wb') as f:
                f.write(image_bytes)

            logger.debug(f"图片保存成功: {filepath}")
            return filepath
        except Exception as e:
            logger.error(f"图片保存失败: {e}")
            return None

    @staticmethod
    def _content_type_to_ext(content_type: str) -> str:
        mapping = {
            "image/png": "png",
            "image/jpeg": "jpg",
            "image/jpg": "jpg",
            "image/gif": "gif",
            "image/bmp": "bmp",
            "image/tiff": "tiff",
            "image/webp": "webp",
            "image/svg+xml": "svg",
        }
        return mapping.get(content_type, "png")


image_extractor = ImageExtractor()
