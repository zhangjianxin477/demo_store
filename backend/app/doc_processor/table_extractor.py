import os
import re
import io
import logging
from typing import List, Dict, Any, Optional, Tuple

from app.core.config import settings

logger = logging.getLogger(__name__)


class TableExtractor:
    def extract_from_pdf(self, file_path: str) -> List[Dict[str, Any]]:
        tables = []
        try:
            import fitz
            tables = self._extract_with_fitz(file_path)
        except ImportError:
            pass

        if not tables:
            try:
                import pdfplumber
                tables = self._extract_with_pdfplumber(file_path)
            except ImportError:
                pass

        if not tables:
            try:
                import camelot
                tables = self._extract_with_camelot(file_path)
            except ImportError:
                pass

        return tables

    def extract_from_docx(self, content: bytes) -> List[Dict[str, Any]]:
        tables = []
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))
            for table_idx, table in enumerate(doc.tables):
                rows = []
                for row in table.rows:
                    cells = []
                    for cell in row.cells:
                        cell_text = cell.text.strip().replace('\n', ' ')
                        cells.append(cell_text)
                    rows.append(cells)
                if rows:
                    md = self._rows_to_markdown(rows)
                    tables.append({
                        "index": table_idx,
                        "rows": len(rows),
                        "cols": len(rows[0]) if rows else 0,
                        "markdown": md,
                        "has_header": True,
                    })
        except Exception as e:
            logger.error(f"DOCX 表格提取失败: {e}")
        return tables

    def extract_from_xlsx(self, content: bytes) -> List[Dict[str, Any]]:
        tables = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() if c is not None else "" for c in row]
                    if any(cells):
                        rows.append(cells)
                if rows:
                    md = self._rows_to_markdown(rows)
                    tables.append({
                        "sheet": sheet_name,
                        "rows": len(rows),
                        "cols": len(rows[0]) if rows else 0,
                        "markdown": md,
                        "has_header": True,
                    })
            wb.close()
        except Exception as e:
            logger.error(f"XLSX 表格提取失败: {e}")
        return tables

    def extract_from_csv(self, content: bytes) -> List[Dict[str, Any]]:
        import csv
        tables = []
        for encoding in ['utf-8', 'utf-8-sig', 'gbk', 'gb2312', 'latin-1']:
            try:
                text = content.decode(encoding)
                break
            except (UnicodeDecodeError, UnicodeError):
                continue
        else:
            text = content.decode('utf-8', errors='replace')

        reader = csv.reader(io.StringIO(text))
        rows = []
        for row in reader:
            cells = [c.strip() for c in row]
            if any(cells):
                rows.append(cells)

        if rows:
            md = self._rows_to_markdown(rows)
            tables.append({
                "rows": len(rows),
                "cols": len(rows[0]) if rows else 0,
                "markdown": md,
                "has_header": True,
            })
        return tables

    def extract_from_markdown(self, text: str) -> List[Dict[str, Any]]:
        tables = []
        table_pattern = re.compile(
            r'((?:^\|.+\|$\n?)+)',
            re.MULTILINE
        )
        for match in table_pattern.finditer(text):
            table_text = match.group(1).strip()
            rows = []
            for line in table_text.split('\n'):
                line = line.strip()
                if not line:
                    continue
                if re.match(r'^\|[\s\-:]+\|$', line):
                    continue
                cells = [c.strip() for c in line.split('|')[1:-1]]
                if cells:
                    rows.append(cells)
            if rows:
                md = self._rows_to_markdown(rows)
                tables.append({
                    "rows": len(rows),
                    "cols": len(rows[0]) if rows else 0,
                    "markdown": md,
                    "has_header": True,
                })
        return tables

    def extract_from_html(self, content: bytes) -> List[Dict[str, Any]]:
        tables = []
        try:
            from bs4 import BeautifulSoup
            for encoding in ['utf-8', 'gbk', 'latin-1']:
                try:
                    text = content.decode(encoding)
                    break
                except (UnicodeDecodeError, UnicodeError):
                    continue
            else:
                text = content.decode('utf-8', errors='replace')

            soup = BeautifulSoup(text, 'html.parser')
            for table_idx, table in enumerate(soup.find_all('table')):
                rows = []
                for tr in table.find_all('tr'):
                    cells = []
                    for td in tr.find_all(['td', 'th']):
                        cells.append(td.get_text(strip=True))
                    if cells:
                        rows.append(cells)
                if rows:
                    md = self._rows_to_markdown(rows)
                    tables.append({
                        "index": table_idx,
                        "rows": len(rows),
                        "cols": len(rows[0]) if rows else 0,
                        "markdown": md,
                        "has_header": True,
                    })
        except Exception as e:
            logger.error(f"HTML 表格提取失败: {e}")
        return tables

    def _extract_with_fitz(self, file_path: str) -> List[Dict[str, Any]]:
        import fitz
        tables = []
        doc = fitz.open(file_path)
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_tables = page.find_tables()
            for table_idx, table in enumerate(page_tables):
                table_data = table.extract()
                if table_data and len(table_data) > 0:
                    md = self._rows_to_markdown(table_data)
                    tables.append({
                        "page": page_idx + 1,
                        "index": table_idx,
                        "rows": len(table_data),
                        "cols": len(table_data[0]) if table_data else 0,
                        "markdown": md,
                        "bbox": table.bbox,
                        "has_header": True,
                    })
        doc.close()
        return tables

    def _extract_with_pdfplumber(self, file_path: str) -> List[Dict[str, Any]]:
        import pdfplumber
        tables = []
        with pdfplumber.open(file_path) as pdf:
            for page_idx, page in enumerate(pdf.pages):
                page_tables = page.extract_tables()
                for table_idx, table_data in enumerate(page_tables):
                    if table_data and len(table_data) > 0:
                        md = self._rows_to_markdown(table_data)
                        tables.append({
                            "page": page_idx + 1,
                            "index": table_idx,
                            "rows": len(table_data),
                            "cols": len(table_data[0]) if table_data else 0,
                            "markdown": md,
                            "has_header": True,
                        })
        return tables

    def _extract_with_camelot(self, file_path: str) -> List[Dict[str, Any]]:
        import camelot
        tables = []
        try:
            camelot_tables = camelot.read_pdf(file_path, pages='all', flavor='lattice')
        except Exception:
            try:
                camelot_tables = camelot.read_pdf(file_path, pages='all', flavor='stream')
            except Exception as e:
                logger.warning(f"Camelot 表格提取失败: {e}")
                return tables

        for table_idx, table in enumerate(camelot_tables):
            df = table.df
            if df.empty:
                continue
            rows = []
            for _, row in df.iterrows():
                cells = [str(v).strip() for v in row.values]
                rows.append(cells)
            if rows:
                md = self._rows_to_markdown(rows)
                tables.append({
                    "page": table.page,
                    "index": table_idx,
                    "rows": len(rows),
                    "cols": len(rows[0]) if rows else 0,
                    "markdown": md,
                    "has_header": True,
                })
        return tables

    @staticmethod
    def _rows_to_markdown(rows: List[List[str]]) -> str:
        if not rows:
            return ""
        max_cols = max(len(row) for row in rows)
        normalized = []
        for row in rows:
            padded = [str(cell).strip() if cell else "" for cell in row]
            while len(padded) < max_cols:
                padded.append("")
            normalized.append(padded)

        header = normalized[0]
        lines = ["| " + " | ".join(header) + " |"]
        lines.append("| " + " | ".join(["---"] * max_cols) + " |")
        for row in normalized[1:]:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)


table_extractor = TableExtractor()
